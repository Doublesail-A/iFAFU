"""Publish the school's arrangements as the original iFAFU Holiday records.

Only public documents are read. No student timetable or login is involved.
The APK consumes the resulting small JSON; document readers run here, in CI.
"""
import argparse
import io
import json
import re
import ssl
import sys
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen
from urllib.error import HTTPError

import pdfplumber
import openpyxl
import xlrd
from bs4 import BeautifulSoup
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / 'data/school-calendars.json'
NAMES = ['元旦', '春节', '清明节', '劳动节', '端午节', '中秋节', '国庆节', '寒假', '暑假', '校运会', '运动会']
LABELS = '|'.join(NAMES)
RANGE = re.compile(r'(?:(\d{4})年)?(\d{1,2})月(\d{1,2})日?(?:至|-|—|～|~)(?:(\d{1,2})月)?(\d{1,2})日?')
DAY = re.compile(r'(?:(\d{4})年)?(\d{1,2})月(\d{1,2})日')
MAKEUP = re.compile(r'(\d{1,2})月(\d{1,2})日上(\d{1,2})月(\d{1,2})日的?(?:课课程|课程|课)')


def school_url(url):
    parsed = urlparse(url)
    if parsed.scheme != 'https' or not parsed.hostname.endswith('.fafu.edu.cn'):
        raise ValueError('Unexpected school document origin')
    return url


def fetch(url):
    request = Request(url, headers={'User-Agent': 'iFAFU-calendar-sync/1.0'})
    # The school also uses RSA TLS suites. Keep normal certificate verification.
    context = ssl.create_default_context()
    context.set_ciphers('DEFAULT')
    with urlopen(request, timeout=25, context=context) as response:
        body = response.read(12 * 1024 * 1024 + 1)
        if not body or len(body) > 12 * 1024 * 1024:
            raise ValueError('Invalid document size')
        return body


def document_text(data, url):
    ext = Path(urlparse(url).path).suffix.lower()
    if ext == '.pdf':
        with pdfplumber.open(io.BytesIO(data)) as document:
            result = []
            for page in document.pages:
                saturday = min((c for c in page.chars if c['text'] == '六'), key=lambda c: c['top'])
                friday = next(c for c in page.chars if c['text'] == '五' and abs(c['top'] - saturday['top']) < 3 and c['x0'] < saturday['x0'])
                center = (saturday['x0'] + saturday['x1']) / 2
                previous = (friday['x0'] + friday['x1']) / 2
                result.append(page.crop((center + (center - previous) / 2, 0, page.width, page.height)).extract_text() or '')
            return '\n'.join(result)
    if ext == '.docx':
        return '\n'.join(row.cells[-1].text for table in Document(io.BytesIO(data)).tables for row in table.rows)
    if ext in {'.xls', '.xlsx'}:
        if ext == '.xls':
            workbook = xlrd.open_workbook(file_contents=data)
            sheets = [[[sheet.cell_value(y, x) for x in range(sheet.ncols)] for y in range(sheet.nrows)] for sheet in workbook.sheets()]
        else:
            workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                sheets = [list(sheet.values) for sheet in workbook.worksheets]
            finally:
                workbook.close()
        result = []
        for rows in sheets:
            columns = [x for row in rows for x, value in enumerate(row) if '备注' in str(value).replace(' ', '')]
            if not columns:
                raise ValueError('No remark column')
            result.extend(str(row[columns[0]] or '') for row in rows if len(row) > columns[0])
        return '\n'.join(result)
    content = BeautifulSoup(data, 'html.parser').select_one('.wp_articlecontent')
    if content is None:
        raise ValueError('No calendar content')
    return content.get_text()


def parse(text, year, term, source=''):
    first, last = map(int, year.split('-'))
    if last != first + 1 or term not in ('1', '2'):
        raise ValueError('Invalid semester')
    text = re.sub(r'\s+', '', text)
    text = re.sub(r'[（(](?:周|星期)[一二三四五六日天][）)]', '', text).replace('－', '-').replace('：', ':')

    def as_date(month, day, explicit=None):
        return date(int(explicit) if explicit else (first if term == '1' and int(month) >= 8 else last), int(month), int(day))

    opening = re.search(r'(\d{1,2})月(\d{1,2})日:?(?:学生|老生|全体学生)(?:正式)?上课', text)
    if not opening:
        raise ValueError('No first teaching day')
    start = as_date(*opening.groups())
    first_week = start - timedelta(days=(start.weekday() + 1) % 7)
    closed, moves, spans = {}, {}, []

    def name_at(match):
        suffix = re.split(r'\d', text[match.end():match.end() + 24])[0]
        following = re.match(rf'^[,:]?({LABELS})', suffix)
        if following:
            return following[1]
        preceding = re.search(rf'((?:{LABELS})(?:[、,](?:{LABELS}))*)(?:时间)?[:]?$' , text[max(0, match.start() - 30):match.start()])
        return preceding[1] if preceding else None

    def close(begin, end, name):
        if not 0 <= (end - begin).days <= 100:
            raise ValueError('Invalid closure range')
        for offset in range((end - begin).days + 1):
            closed[(begin + timedelta(days=offset)).isoformat()] = name

    for match in RANGE.finditer(text):
        name = name_at(match)
        if not name:
            continue
        explicit, month, day, end_month, end_day = match.groups()
        begin = as_date(month, day, explicit)
        end = as_date(end_month or month, end_day, explicit)
        if end < begin and int(end_month or month) < begin.month:
            end = end.replace(year=end.year + 1)
        close(begin, end, name)
        spans.append((match.start(), match.end()))
    for match in DAY.finditer(text):
        if any(begin <= match.start() < end for begin, end in spans):
            continue
        name = name_at(match)
        suffix = text[match.end():match.end() + 30]
        if not name or '放假' not in suffix[:16]:
            continue
        explicit, month, day = match.groups()
        begin = end = as_date(month, day, explicit)
        if '与周末连休' in suffix:
            if begin.weekday() == 0:
                begin -= timedelta(days=2)
            elif end.weekday() == 4:
                end += timedelta(days=2)
        close(begin, end, name)
    for match in MAKEUP.finditer(text):
        destination = as_date(*match.groups()[:2]).isoformat()
        original = as_date(*match.groups()[2:]).isoformat()
        if original in moves and moves[original] != destination:
            raise ValueError('Conflicting makeup')
        moves[original] = destination
    if len(re.findall(r'的(?:课课程|课程|课)', text)) != len(list(MAKEUP.finditer(text))):
        raise ValueError('Unrecognized makeup expression')
    deferred = set()
    for match in re.finditer(rf'((?:{LABELS})(?:[、,](?:{LABELS}))*)放假安排以国务院[^。;]*通知为准', text):
        deferred.update(name for name in NAMES if name in match[1])
    deferred -= {label for value in closed.values() for label in re.split('[、,]', value)}
    for name in NAMES:
        if re.search(rf'{name}[^。;]{{0,60}}放假', text) and name not in deferred and not any(name in value for value in closed.values()):
            raise ValueError('Unrecognized holiday expression')
    if not closed and not deferred:
        raise ValueError('No holiday arrangements')
    return {'school': 'FAFU', 'year': year, 'term': term, 'firstWeek': first_week.isoformat(), 'closed': closed, 'moves': moves, 'deferred': sorted(deferred), 'source': source}


def holiday_records(closed, moves):
    """Same from/days/changes model used by the upstream application."""
    result = []
    for iso, name in sorted(closed.items()):
        if result and result[-1]['name'] == name and (date.fromisoformat(iso) - date.fromisoformat(result[-1]['from'])).days == result[-1]['days']:
            result[-1]['days'] += 1
        else:
            result.append({'name': name, 'from': iso, 'days': 1, 'changes': {}})
    if moves:
        result.append({'name': '调课', 'from': min(moves), 'days': 0, 'changes': dict(sorted(moves.items()))})
    return result


def catalog():
    queue, visited, result = ['https://jwc.fafu.edu.cn/2011/list.htm'], set(), {}
    while queue and len(visited) < 20:
        url = queue.pop(0)
        if url in visited:
            continue
        visited.add(url)
        soup = BeautifulSoup(fetch(school_url(url)), 'html.parser')
        for link in soup.select('a[href]'):
            title = re.sub(r'\s+', '', link.get('title') or link.get_text()).replace('－', '-').replace('—', '-')
            year = re.search(r'\d{4}-\d{4}', title)
            term = re.search(r'第([12一二])学期', title)
            href = urljoin(url, link['href'])
            if '校历' in title and year and term:
                number = {'一': '1', '二': '2'}.get(term[1], term[1])
                result.setdefault((year[0], number), school_url(href))
            if re.fullmatch(r'https://jwc\.fafu\.edu\.cn/2011/list\d*\.htm', href) and href not in visited:
                queue.append(href)
    if not result:
        raise ValueError('Empty school calendar directory')
    return result


def sync(output=OUTPUT):
    previous = json.loads(output.read_text('utf-8'))['calendars'] if output.exists() else []
    entries = {(item['year'], item['term']): item for item in previous}
    failures, national = [], {}
    for (year, term), source in sorted(catalog().items()):
        try:
            data = fetch(source)
            document = source
            if Path(urlparse(source).path).suffix.lower() not in {'.pdf', '.docx', '.xls', '.xlsx'}:
                soup = BeautifulSoup(data, 'html.parser')
                content = soup.select_one('.wp_articlecontent')
                attachment = content.select_one('[pdfsrc]') if content else None
                if attachment:
                    document = urljoin(source, attachment['pdfsrc'])
                elif content:
                    document = next((urljoin(source, a['href']) for a in content.select('a[href]') if Path(urlparse(a['href']).path).suffix.lower() in {'.pdf', '.docx', '.xls', '.xlsx'}), source)
                if document != source:
                    data = fetch(school_url(document))
            value = parse(document_text(data, document), year, term, source)
            # Only supplement holidays that the school explicitly defers to the
            # State Council. Workdays never imply a teaching weekday.
            if value['deferred']:
                for y in map(int, year.split('-')):
                    if y not in national:
                        try:
                            national[y] = json.loads(fetch(f'https://raw.githubusercontent.com/NateScarlet/holiday-cn/master/{y}.json'))
                        except HTTPError as error:
                            if error.code != 404:
                                raise
                            national[y] = {'days': [], 'papers': []}  # announcement not published yet
                    announcement = national[y]
                    if announcement.get('days') and not any(urlparse(p).hostname.endswith('gov.cn') for p in announcement.get('papers', [])):
                        raise ValueError('No State Council provenance')
                    for day in announcement.get('days', []):
                        when = date.fromisoformat(day['date'])
                        if day['name'] in value['deferred'] and day['isOffDay'] and date.fromisoformat(value['firstWeek']) <= when < date.fromisoformat(value['firstWeek']) + timedelta(weeks=32):
                            value['closed'][day['date']] = day['name']
            item = {key: value[key] for key in ('school', 'year', 'term', 'firstWeek', 'source')}
            item['holidays'] = holiday_records(value['closed'], value['moves'])
            entries[year, term] = item
        except Exception as error:
            # Preserve previously published valid records; an unknown document
            # must not replace them with an empty list.
            message = f'{year}/{term}: {error}'
            if int(year.split('-')[1]) >= date.today().year:
                failures.append(message)
            else:
                print('Historical document skipped: ' + message, file=sys.stderr)
    payload = {'version': 1, 'checkedOn': date.today().isoformat(), 'calendars': [entries[key] for key in sorted(entries)]}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    for error in failures:
        print(error, file=sys.stderr)
    if failures:
        raise RuntimeError('Some school documents could not be updated')
    print(f'Published {len(entries)} semester records')


if __name__ == '__main__':
    arguments = argparse.ArgumentParser()
    arguments.add_argument('--output', type=Path, default=OUTPUT)
    sync(arguments.parse_args().output)
